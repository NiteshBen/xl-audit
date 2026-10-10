"""Tests for xl_audit — each check exercised on small generated workbooks."""
import json

import pytest
from openpyxl import Workbook

from xl_audit import analyze, main


def make_book(tmp_path, name="t.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    path = tmp_path / name
    return wb, ws, path


def save(ws_wb, path):
    ws_wb[0].save(path)
    return path


def findings_for(path):
    findings, summary = analyze(path)
    return findings, summary


def test_clean_workbook_has_no_findings(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Item", "Amount"])
    for i, amount in enumerate([10, 20, 30, 40], start=2):
        ws.append([f"item {i}", amount])
    ws.cell(row=7, column=2, value="=SUM(B2:B6)")     # empty gap below range: fine
    findings, summary = findings_for(save((wb, ws), path))
    assert not [f for f in findings if f.severity in ("error", "warn")]
    assert summary["formulas"] == 1


def test_sum_scope_catches_missed_row(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Item", "Amount"])
    for i, amount in enumerate([10, 20, 30, 40], start=2):   # data B2:B5
        ws.append([f"item {i}", amount])
    ws.cell(row=7, column=2, value="=SUM(B2:B4)")    # stops one row short of B5
    findings, _ = findings_for(save((wb, ws), path))
    assert any(f.check == "sum-scope" and "B5" in f.title for f in findings)


def test_sum_scope_tolerates_subtotal_below_range(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Item", "Amount"])
    for i, amount in enumerate([10, 20, 30], start=2):
        ws.append([f"item {i}", amount])
    ws.cell(row=5, column=2, value="=SUM(B2:B4)")    # exact range
    ws.cell(row=6, column=2, value="=B5*2")          # formula below: not missed data
    findings, _ = findings_for(save((wb, ws), path))
    assert not any(f.check == "sum-scope" for f in findings)


def test_formula_drift_catches_copy_paste_slip(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Item", "Amount", "Double"])
    for i, amount in enumerate([10, 20, 30, 40], start=2):
        ws.append([f"item {i}", amount])
        ws.cell(row=i, column=3, value=f"=B{i}*2")
    ws.cell(row=4, column=3, value="=B4*3")          # row 4 uses the wrong multiplier
    findings, _ = findings_for(save((wb, ws), path))
    drift = [f for f in findings if f.check == "formula-drift"]
    assert drift and any("C4" in ex for ex in drift[0].examples)


def test_hardcoded_constant_flagged_but_plain_formulas_not(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Amount", "With tax"])
    ws.append([100, "=A2*1.18"])
    ws.append([50, "=A3*1.18"])
    ws.append([25, "=A4+A5"])
    ws.append([5, "=SUM(A2:A5)"])
    ws.append([7, "=A6*2"])                     # benign small-integer multiplier
    findings, _ = findings_for(save((wb, ws), path))
    hard = [f for f in findings if f.check == "hardcoded"]
    assert len(hard) == 1 and len(hard[0].examples) == 2   # grouped per column
    assert all("B" in ex for ex in hard[0].examples)       # the two *1.18 cells
    assert not any("B6" in ex for ex in hard[0].examples)  # *2 not flagged


def test_broken_reference_detected(tmp_path):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Amount"])
    ws.append([100])
    ws.cell(row=2, column=2, value="=#REF!+1")
    findings, _ = findings_for(save((wb, ws), path))
    assert any(f.check == "broken-refs" for f in findings)


def test_cli_end_to_end_exit_codes_and_json(tmp_path, capsys):
    wb, ws, path = make_book(tmp_path)
    ws.append(["Amount"])
    ws.append([100])
    ws.cell(row=3, column=2, value="=#REF!")
    dirty = save((wb, ws), path)

    wb2, ws2, path2 = make_book(tmp_path, name="clean.xlsx")
    ws2.append(["Amount"])
    ws2.append([100])
    ws2.cell(row=3, column=2, value="=A2*2")
    clean = save((wb2, ws2), path2)

    assert main([str(dirty)]) == 1
    assert "xl-audit report" in capsys.readouterr().out
    assert main([str(clean)]) == 0
    assert main([str(tmp_path / "nope.xlsx")]) == 2
    capsys.readouterr()                                   # flush earlier output

    assert main([str(dirty), "--json"]) == 1
    report = json.loads(capsys.readouterr().out)
    assert report["summary"]["formulas"] >= 1
    assert any(f["check"] == "broken-refs" for f in report["findings"])
